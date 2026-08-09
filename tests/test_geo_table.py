"""Geolocation enrichment and the spreadsheet export.

None of these tests reach the network: an address that is private, reserved
or documentation is answered locally, and the hosted lookup is only ever
attempted when a token is configured.
"""
from plaguardsv2.GuardModules import PlagArith, PlagGeo, PlagTable


def test_addresses_are_categorised_before_anything_is_looked_up():
    assert PlagGeo.category("8.8.8.8") == "public"
    assert PlagGeo.category("10.0.0.5") == "private"
    assert PlagGeo.category("127.0.0.1") == "loopback"
    assert PlagGeo.category("169.254.1.1") == "link-local"
    assert PlagGeo.category("not-an-ip") == "invalid"


def test_documentation_ranges_are_not_called_private():
    """ipaddress reports the RFC 5737 blocks as private, which would label a
    documentation address as somebody's internal network."""
    for value in ("192.0.2.1", "198.51.100.42", "203.0.113.7"):
        assert PlagGeo.category(value) == "documentation"


def test_a_non_public_address_is_answered_without_a_lookup():
    record = PlagGeo.locate("10.0.0.5")
    assert record["category"] == "private"
    assert "no public geolocation" in record["note"]
    # Every field is present, so callers never have to guard for a missing key.
    assert set(PlagGeo.FIELDS).issubset(record)


def test_a_public_address_says_so_when_nothing_is_configured():
    record = PlagGeo.locate("8.8.8.8", conn=None)
    assert record["category"] == "public"
    assert record["note"] == "no geolocation source configured"


def test_enrich_covers_each_address_once_and_strips_the_port():
    class Finding:
        def __init__(self, type, value):
            self.type, self.value = type, value

    findings = [
        Finding("ip", "10.0.0.5"),
        Finding("ip", "10.0.0.5"),
        Finding("ip_port", "198.51.100.42:4444"),
        Finding("domain", "example.invalid"),
    ]
    located = PlagGeo.enrich(findings, conn=None)
    assert sorted(located) == ["10.0.0.5", "198.51.100.42"]


def _analysis():
    return {"findings": [
        {"severity": "high", "type": "ip", "value": "198.51.100.42",
         "mitre": ["T1071"], "description": "C2 address",
         "intel": {"virustotal": {"verdict": "clean", "detail": "0 malicious"},
                   "abuseipdb": {"verdict": "malicious", "detail": "90%"}},
         "triage_status": "confirmed_threat", "context": "IP=198.51.100.42:4444"},
        {"severity": "info", "type": "domain", "value": "malicious.example.invalid",
         "mitre": [], "description": "", "intel": {},
         "triage_status": "unreviewed", "context": "C2="},
    ]}


def test_rows_carry_the_findings_and_their_geolocation():
    geo = {"198.51.100.42": PlagGeo.locate("198.51.100.42")}
    rows = PlagTable.build_rows(_analysis(), geo)

    assert [r["n"] for r in rows] == [1, 2]
    assert rows[0]["defanged_value"] == "198[.]51[.]100[.]42"
    assert rows[0]["geo_category"] == "documentation"
    # The worst provider verdict is the headline.
    assert rows[0]["intel_verdict"] == "malicious"
    assert "abuseipdb" in rows[0]["intel_detail"]
    # A finding that is not an address simply has empty geolocation columns.
    assert rows[1]["geo_category"] == ""
    assert rows[1]["intel_verdict"] == "not checked"


def test_csv_has_a_header_for_every_column_and_opens_in_excel():
    rows = PlagTable.build_rows(_analysis(), {})
    data = PlagTable.to_csv(rows)
    # A BOM, or Excel mangles anything non-ASCII.
    assert data.startswith(b"\xef\xbb\xbf")
    lines = data.decode("utf-8-sig").splitlines()
    assert lines[0].count(",") == len(PlagTable.COLUMNS) - 1
    assert len(lines) == len(rows) + 1


def test_xlsx_is_a_real_workbook():
    if not PlagTable.xlsx_available():
        return
    rows = PlagTable.build_rows(_analysis(), {})
    data = PlagTable.to_xlsx(rows)
    assert data[:2] == b"PK"  # xlsx is a zip

    import io

    import openpyxl
    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    assert sheet.max_row == len(rows) + 1
    assert sheet.max_column == len(PlagTable.COLUMNS)
    assert sheet.freeze_panes == "A2"


def test_arithmetic_folds_only_arithmetic():
    assert PlagArith.evaluate("108+1") == 109
    assert PlagArith.evaluate("0x6d") == 109
    assert PlagArith.evaluate("&H6D") == 109  # VBScript hex
    assert PlagArith.evaluate("109 ^ 0") == 109
    assert PlagArith.to_char("50+2") == "4"
    # Anything that is not a number or an operator is refused outright.
    for bad in ("x+1", "__import__('os')", "open('f')", "", "a" * 300):
        assert PlagArith.evaluate(bad) is None
