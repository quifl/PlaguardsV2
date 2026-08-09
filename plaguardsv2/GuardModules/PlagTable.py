"""Findings as a spreadsheet.

A PDF is for reading; a table is for pivoting, sorting and pasting into a
ticket. Same data as the report - every finding with its severity, MITRE
mapping, threat-intel verdict and analyst decision - plus the geolocation
columns for any address, in the spirit of HolmesGeo's enriched output.

CSV is written with the standard library. XLSX needs openpyxl and is simply
unavailable if it is not installed, rather than being a hard dependency.
"""
from __future__ import annotations

import csv
import io

from . import PlagGeo, PlagGrep, PlagMitre

COLUMNS = [
    ("n", "#"),
    ("severity", "Severity"),
    ("type", "Type"),
    ("value", "Value"),
    ("defanged_value", "Value (defanged)"),
    ("mitre", "MITRE ATT&CK"),
    ("description", "What it means"),
    ("intel_verdict", "Threat intel"),
    ("intel_detail", "Threat intel detail"),
    ("triage_status", "Analyst verdict"),
    ("context", "Seen in context"),
    # HolmesGeo-style enrichment, populated for addresses only.
    ("geo_category", "IP category"),
    ("geo_city", "City"),
    ("geo_latitude", "City latitude"),
    ("geo_longitude", "City longitude"),
    ("geo_country", "Country"),
    ("geo_country_code", "Country code"),
    ("geo_continent", "Continent"),
    ("geo_asn", "ASN number"),
    ("geo_organization", "ASN organization"),
    ("geo_network", "Network"),
    ("geo_reverse_dns", "Reverse DNS"),
]

MAX_CELL = 2000


def _intel_summary(intel: dict) -> tuple[str, str]:
    """Worst verdict across providers, and who said what."""
    ranked = ["malicious", "suspicious", "unknown", "error", "clean"]
    scored = {
        source: result for source, result in (intel or {}).items()
        if source != "scope" and result.get("verdict") != "not_checked"
    }
    if not scored:
        return "not checked", ""
    verdicts = [r.get("verdict", "unknown") for r in scored.values()]
    headline = next((v for v in ranked if v in verdicts), "unknown")
    detail = "; ".join(
        f"{source}: {result.get('verdict', '')}"
        + (f" ({result['detail']})" if result.get("detail") else "")
        for source, result in sorted(scored.items())
    )
    return headline.replace("_", " "), detail


def build_rows(analysis: dict, geo: dict | None = None) -> list[dict]:
    """One row per finding, with geolocation filled in for addresses."""
    geo = geo or {}
    rows = []
    for index, finding in enumerate(analysis.get("findings") or [], 1):
        get = finding.get if isinstance(finding, dict) else lambda k, d=None: getattr(finding, k, d)
        value = str(get("value", "") or "")
        ftype = get("type", "") or ""
        verdict, detail = _intel_summary(get("intel", {}) or {})

        # Every geolocation column exists on every row, even for a finding
        # that is not an address, so the table stays rectangular.
        row = {f"geo_{field}": "" for field in PlagGeo.FIELDS if field != "ip"}
        row.update({
            "n": index,
            "severity": get("severity", "") or "",
            "type": ftype,
            "value": value,
            "defanged_value": PlagGrep.defang(value, ftype),
            "mitre": "; ".join(PlagMitre.describe(t) for t in (get("mitre") or [])),
            "description": get("description", "") or "",
            "intel_verdict": verdict,
            "intel_detail": detail,
            "triage_status": str(get("triage_status", "unreviewed")).replace("_", " "),
            "context": (get("context", "") or "")[:MAX_CELL],
        })

        host = value.rsplit(":", 1)[0] if ftype == "ip_port" else value
        record = geo.get(host)
        if record:
            for field in PlagGeo.FIELDS:
                if field == "ip":
                    continue
                row[f"geo_{field}"] = record.get(field, "")
        rows.append(row)
    return rows


def to_csv(rows: list[dict]) -> bytes:
    buf = io.StringIO(newline="")
    writer = csv.writer(buf)
    writer.writerow([label for _key, label in COLUMNS])
    for row in rows:
        writer.writerow([row.get(key, "") for key, _label in COLUMNS])
    # BOM so Excel opens a UTF-8 CSV without mangling non-ASCII values.
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def xlsx_available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def to_xlsx(rows: list[dict], title: str = "Findings") -> bytes:
    """The same table as a real spreadsheet - frozen header, filters, widths."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title[:31] or "Findings"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="0D1A2B")
    for column, (_key, label) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=column, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    severity_fill = {
        "high": PatternFill("solid", start_color="F6DBE0"),
        "medium": PatternFill("solid", start_color="FBEED3"),
        "low": PatternFill("solid", start_color="D9EFE2"),
        "info": PatternFill("solid", start_color="E6E9F0"),
    }
    for r, row in enumerate(rows, 2):
        for c, (key, _label) in enumerate(COLUMNS, 1):
            cell = sheet.cell(row=r, column=c, value=row.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=(key == "context"))
        fill = severity_fill.get(str(row.get("severity", "")).lower())
        if fill:
            sheet.cell(row=r, column=2).fill = fill

    widths = {"value": 38, "defanged_value": 38, "mitre": 34, "description": 44,
              "intel_detail": 44, "context": 52}
    for column, (key, label) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = widths.get(key, max(12, len(label) + 2))

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(1, len(rows) + 1)}"

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
